import re
import string
from datetime import UTC, datetime

import psycopg2
from psycopg2.extensions import ISOLATION_LEVEL_AUTOCOMMIT
from sqlalchemy import create_engine

from ust.python.util import config
from ust.python.util.logger_factory import error_logger, logger


def _require_config(*names):
    missing = [name for name in names if not getattr(config, name, None)]
    if missing:
        raise RuntimeError(
            "Missing database config values: "
            + ", ".join(missing)
            + ". Create ust/python/util/config.py or set the matching UST_* environment variables."
        )


def connect_db(db_name=None, schema='public'):
    _require_config('db_ip', 'db_user', 'db_password')
    db_name = db_name or getattr(config, 'db_name', None)
    if not db_name:
        _require_config('db_name')
    try:
        options = f'-csearch_path="{schema}"'
        conn = psycopg2.connect(
                    host=config.db_ip,
                    user=config.db_user,
                    password=config.db_password,
                    dbname=db_name,
                    keepalives=1,
                    keepalives_idle=30,
                    keepalives_interval=10,
                    keepalives_count=5,
                    options=options)
    except Exception as e:
        logger.error('Unable to connect to database %s: %s', db_name, e)
        raise

    conn.set_isolation_level(ISOLATION_LEVEL_AUTOCOMMIT)
    return conn


def get_engine(db_name=None, schema='public'):
    _require_config('db_connection_string')
    db_name = db_name or getattr(config, 'db_name', None)
    if not db_name:
        _require_config('db_name')
    try:
        engine = create_engine(config.db_connection_string + db_name, connect_args={'options': f'-csearch_path="{schema}"'})
        return engine
    except Exception as e:
        logger.error('Error creating database engine: %s', e)
        raise


def get_sqlserver_connection_string(host, db, user, passwd, driver = 'ODBC Driver 18 for SQL Server'):
    return f"DRIVER={{{driver}}};SERVER={host};DATABASE={db};uid={user};pwd={passwd};TrustServerCertificate=yes;Encrypt=no;"


def connect_sqlserver_db(host, db, user, passwd, driver='ODBC Driver 18 for SQL Server'):
    import pyodbc

    conn_str = get_sqlserver_connection_string(host, db, user, passwd, driver=driver)
    conn =  pyodbc.connect(conn_str) 
    logger.info('Connected to %s', db)
    return conn


def get_sqlserver_engine(host, db, user, passwd, driver='ODBC Driver 18 for SQL Server'):
    conn_str = get_sqlserver_connection_string(host, db, user, passwd, driver=driver)
    return create_engine(f"mssql+pyodbc:///?odbc_connect={conn_str}", fast_executemany=True)    


def get_view_sql(view_name):
    conn = connect_db()
    cur = conn.cursor()
    sql = f"select pg_get_viewdef('{view_name}', true)"
    process_sql(conn, cur, sql)
    view_sql =  cur.fetchone()[0]
    cur.close()
    conn.close()
    return view_sql

 
def process_view_sql(sql):
    i = sql.find('FROM')
    to_sql = sql[:i].replace('SELECT DISTINCT ','').replace('SELECT ','')
    to_sql = to_sql.replace('SELECT DISTINCT\n','').replace('SELECT\n','')
    i2 = sql.find('ORDER BY ')
    from_sql = sql[i:i2]
    return (to_sql, from_sql)


def extract_col_alias(col_def):
    i = col_def.find('.') + 1
    alias = col_def[i:].replace('"','')
    return alias


def process_col_sql(sql):
    i = sql.find(' AS ')
    if i == -1:
        i = len(sql)
    select_col = sql[:i]
    if ' AS ' in sql:
        alias = sql[i:].replace(' AS ','').replace('"','')
    else:
        alias = extract_col_alias(select_col)
    return alias, select_col


def get_select_cols(sql):
    orig_cols = re.split(r',\s*(?![^()]*\))', sql)
    cols = {}
    for col in orig_cols:
        column = col.strip()
        c = process_col_sql(column)
        alias = c[0]
        col_def = c[1]
        cols[alias] = col_def
    return cols


def autowidth(worksheet):
    for col in worksheet.columns:
         autowidth_column(worksheet, col)


def autowidth_column(worksheet, column):
    max_length = 0
    column_letter = column[0].column_letter # get the column name
    for cell in column:
        value = cell.value
        if value is None:
            continue
        max_length = max(max_length, len(str(value)))
        adjusted_width = (max_length + 2) * 1.2
        worksheet.column_dimensions[column_letter].width = adjusted_width    

def add_ws_filter(worksheet):
    worksheet.auto_filter.ref = worksheet.dimensions
    return worksheet


def get_today_string():
    return datetime.now(tz=UTC).strftime('%Y-%m-%d')


def get_control_id(ust_or_release, organization_id):
    ust_or_release = ust_or_release.lower()
    conn = connect_db()
    cur = conn.cursor()

    sql = f"select count(*) from public.{ust_or_release}_control where organization_id = %s"
    process_sql(conn, cur, sql, params=(organization_id, ))
    cnt = cur.fetchone()[0]
    if cnt == 0:
        logger.error('No data in %s_control; unable to proceed.',ust_or_release)
        cur.close()
        conn.close()
        raise LookupError(f'No rows found in public.{ust_or_release}_control for organization_id {organization_id!r}.')
    sql = f"select max({ust_or_release}_control_id) from public.{ust_or_release}_control where organization_id = %s"
    process_sql(conn, cur, sql, params=(organization_id, ))
    control_id = cur.fetchone()[0]

    if cnt > 1:
        logger.info('Found multiple Control IDs in %s_control; using newest one (%s)', ust_or_release, str(control_id) )
    else:
        logger.info('Setting Control ID to %s', str(control_id))

    cur.close()
    conn.close()

    return control_id


def get_table_existence(table_name, table_schema='public'):
    conn = connect_db()
    cur = conn.cursor()
    sql = """select count(*) from information_schema.tables 
              where table_schema = %s and table_name = %s"""
    process_sql(conn, cur, sql, params=(table_schema, table_name))
    cnt = cur.fetchone()[0]
    cur.close()
    conn.close()
    return cnt != 0


def get_selenium_driver(url):
    import time

    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    from selenium.webdriver.chrome.service import Service

    options = Options()
    service = Service()
    # comment out the line below to show the browser (use only for debugging)
    options.add_argument('--headless')
    # uncomment the line below to keep the browser window open; use only for debugging
    # options.add_experimental_option('detach', True)
    driver = webdriver.Chrome(service=service, options=options)
    driver.get(url)
    while driver.execute_script("return document.readyState") != "complete":
        time.sleep(1)
    time.sleep(5)
    return driver


def get_org_from_control_id(control_id, ust_or_release):
    if ust_or_release.lower() not in ['ust','release']:
        logger.error('Invalid value %s for ust_or_release. Valid values are ust or release. Exiting...', ust_or_release)
        raise ValueError(f"Invalid ust_or_release {ust_or_release!r}; expected 'ust' or 'release'.")
    conn = connect_db()
    cur = conn.cursor()    
    table_name = ust_or_release.lower() + '_control'
    column_name = table_name + '_id'
    sql = f"select organization_id from {table_name} where {column_name} = %s"
    process_sql(conn, cur, sql, params=(control_id,))
    ok = True
    org = None 
    try:
        org = cur.fetchone()[0]
    except TypeError:
        logger.warning('No %s with a value of %s found in %s. Exiting', column_name, control_id, table_name)
        ok = False
    cur.close()
    conn.close()
    if ok:
        return org  
    else:
        raise LookupError(f'No {column_name} value {control_id!r} found in {table_name}.')


def delete_all_release_data(control_id):
    conn = connect_db()
    cur = conn.cursor() 

    sql = """delete from public.ust_release_corrective_action_strategy 
            where ust_release_id in (select ust_release_id from ust_release where release_control_id = %s)"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_release_corrective_action_strategy', cur.rowcount)

    sql = """delete from public.ust_release_cause 
            where ust_release_id in (select ust_release_id from ust_release where release_control_id = %s)"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_release_cause', cur.rowcount)

    sql = """delete from public.ust_release_source 
            where ust_release_id in (select ust_release_id from ust_release where release_control_id = %s)"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_release_source', cur.rowcount)

    sql = """delete from public.ust_release_substance 
            where ust_release_id in (select ust_release_id from ust_release where release_control_id = %s)"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_release_substance', cur.rowcount)

    sql = """delete from public.ust_release_geocode 
            where ust_release_id in (select ust_release_id from ust_release where release_control_id = %s)"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_release_geocode', cur.rowcount)

    sql = """delete from public.ust_release where release_control_id = %s"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_release', cur.rowcount)

    conn.commit()
    cur.close()    
    conn.close()


def delete_all_ust_data(control_id):
    conn = connect_db()
    cur = conn.cursor() 

    sql = """delete from public.ust_compartment_dispenser 
             where ust_compartment_id in
                (select ust_compartment_id from public.ust_compartment 
                where ust_tank_id in 
                    (select ust_tank_id from public.ust_tank 
                    where ust_facility_id in
                        (select ust_facility_id from public.ust_facility where ust_control_id = %s)))"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_compartment_dispenser', cur.rowcount)

    sql = """delete from public.ust_piping 
             where ust_compartment_id in
                (select ust_compartment_id from public.ust_compartment 
                where ust_tank_id in 
                    (select ust_tank_id from public.ust_tank 
                    where ust_facility_id in
                        (select ust_facility_id from public.ust_facility where ust_control_id = %s)))"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_piping', cur.rowcount)

    sql = """delete from public.ust_compartment_substance
             where ust_compartment_id in
                (select ust_compartment_id from public.ust_compartment 
                where ust_tank_id in 
                    (select ust_tank_id from public.ust_tank 
                    where ust_facility_id in
                        (select ust_facility_id from public.ust_facility where ust_control_id = %s)))"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_compartment_substance', cur.rowcount)

    sql = """delete from public.ust_compartment
             where ust_tank_id in 
                (select ust_tank_id from public.ust_tank 
                where ust_facility_id in
                     (select ust_facility_id from public.ust_facility where ust_control_id = %s))"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_compartment', cur.rowcount)

    sql = """delete from public.ust_tank_dispenser
             where ust_tank_id in 
                (select ust_tank_id from public.ust_tank 
                where ust_facility_id in
                     (select ust_facility_id from public.ust_facility where ust_control_id = %s))"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_tank_dispenser', cur.rowcount)

    sql = """delete from public.ust_tank_substance 
             where ust_tank_id in 
                (select ust_tank_id from public.ust_tank 
                where ust_facility_id in
                     (select ust_facility_id from public.ust_facility where ust_control_id = %s))"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_tank_substance', cur.rowcount)

    sql = """delete from public.ust_tank 
             where ust_facility_id in
                 (select ust_facility_id from public.ust_facility where ust_control_id = %s)"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_tank', cur.rowcount)

    sql = """delete from public.ust_facility_dispenser
             where ust_facility_id in
                 (select ust_facility_id from public.ust_facility where ust_control_id = %s)"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_facility_dispenser', cur.rowcount)

    sql = """delete from public.ust_facility_geocode
             where ust_facility_id in
                 (select ust_facility_id from public.ust_facility where ust_control_id = %s)"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_facility_geocode', cur.rowcount)

    sql = """delete from public.ust_facility where ust_control_id = %s"""
    process_sql(conn, cur, sql, params=(control_id,))
    logger.info('Deleted %s rows from public.ust_facility', cur.rowcount)

    conn.commit()
    cur.close()    
    conn.close()


def get_schema_from_control_id(control_id, ust_or_release):
    if control_id == 1:
        return None
    org = get_org_from_control_id(control_id, ust_or_release)
    return org.lower() + '_' + ust_or_release.lower()


def get_lookup_tabs(ust_or_release='ust'):
    ust_or_release = ust_or_release.lower()
    if ust_or_release not in ['ust','release']:
        logger.warning('Unknown value passed got get_lookup_tabs(ust_or_release): %s', ust_or_release)
        return 
    conn = connect_db()
    cur = conn.cursor() 
    sql = f"""select table_name, desc_column_name, id_column_name  
            from public.{ust_or_release}_template_lookup_tables
            order by sort_order"""
    process_sql(conn, cur, sql)
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return rows


def get_data_tabs(ust_or_release='ust'):
    ust_or_release = ust_or_release.lower()
    if ust_or_release not in ['ust','release']:
        logger.warning('Unknown value passed got get_data_tabs(ust_or_release): %s', ust_or_release)
        return 
    conn = connect_db()
    cur = conn.cursor() 
    sql = f"""select view_name, template_tab_name  
            from {ust_or_release}_template_data_tables
            order by sort_order"""
    process_sql(conn, cur, sql)
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return rows


def get_headers(table_name, schema='public'):
    conn = connect_db()
    cur = conn.cursor() 
    sql = """select column_name from information_schema.columns
             where table_schema = %s and table_name = %s
             order by ordinal_position"""
    process_sql(conn, cur, sql, params=(schema, table_name))
    headers = [x[0] for x in cur.fetchall()]
    cur.close()
    conn.close()
    return headers


def get_timestamp_str():
    now = datetime.now(tz=UTC)
    return now.strftime('%Y%m%d%H%M%S')


def get_fill_gen(color, color2=None):
    from openpyxl.styles import PatternFill
    if not color2:
        color2 = color
    fill_gen = PatternFill(fill_type='solid',
                       start_color=color,
                       end_color=color2)
    return fill_gen


def get_table_values(table_name, column_name, schema='public'):
    conn = connect_db()
    cur = conn.cursor() 
    sql = f'select distinct "{column_name}" from "{schema}"."{table_name}" order by 1' 
    process_sql(conn, cur, sql)
    values = [r[0] for r in cur.fetchall()]
    cur.close()
    conn.close()    
    return values 


def is_excel(file_path):
    return file_path.lower()[-4:] == 'xlsx' or file_path.lower()[-3:] == 'xls'


def get_pretty_ust_or_release(ust_or_release):
    if ust_or_release.lower() == 'ust':
        return 'UST'
    elif ust_or_release.lower() == 'release':
        return 'Releases'
    else:
        logger.warning('Invalid value for ust_or_release: %s. Valid values are "ust" and "release"')
        return None 


def verify_ust_or_release(ust_or_release):
    if ust_or_release.lower() not in ['ust','release']:
        logger.error("Unknown value '%s' for ust_or_release; valid values are 'ust' and 'release'. Exiting...", ust_or_release)
        raise ValueError(f"Unknown ust_or_release {ust_or_release!r}; expected 'ust' or 'release'.")
    else:
        return ust_or_release.lower()


def get_deagg_table_name(column_name):
    return 'erg_' + column_name.lower().replace(' ','_') + '_deagg'


def get_deagg_datarows_table_name(deagg_table_name):
    return deagg_table_name.replace('_deagg','_datarows_deagg')


def get_epa_region(organization_id):
    conn = connect_db()
    cur = conn.cursor()
    sql = "select epa_region from public.epa_regions where organization_id = %s"
    process_sql(conn, cur, sql, params=(organization_id,))
    try:
        epa_region =  cur.fetchone()[0]
    except TypeError:
        logger.warning('No such organization_id in table public.epa_regions: %s', organization_id)
        epa_region = None 
    cur.close()
    conn.close()
    return epa_region


def get_datatype_sql(data_type, character_maximum_length=None):
    datatype_sql = None 
    if data_type == 'character varying' and not character_maximum_length:
        logger.error('character_maximum_length is required if data_type = character varying')
        raise ValueError('character_maximum_length is required when data_type is character varying.')
    elif character_maximum_length:
        try:
            int(character_maximum_length)
        except (TypeError, ValueError):
            logger.error('character_maximum_length must be an integer (received %s)', character_maximum_length)
            raise ValueError(f'character_maximum_length must be an integer, received {character_maximum_length!r}.')
    if data_type == 'character varying':
        datatype_sql = data_type + '(' + str(character_maximum_length) + ')'
    else:
        datatype_sql = data_type 
    return datatype_sql


def remove_extra_whitespace(string):
    return re.sub(r"(?:(?!\n)\s)+", " ",string)


def get_pretty_query(cursor):
    try:
        query = cursor.query.decode('utf-8')
    except (AttributeError, UnicodeDecodeError) as e:
        logger.warning('Wrote input to cursor.query; unable to print: %s', e)
        return
    query = remove_extra_whitespace(query) + ';'
    return query


def pretty_print_query(cursor):
    print(get_pretty_query(cursor))


def get_element_name_from_colname(column_name):
    return column_name.replace('_',' ').title().replace(' ','')


def comment_every_line(ql_string):
    commented_string = ''
    ql_list = ql_string.split('\n')
    for ql in ql_list:
        if ql == '':
            continue
        commented_string = commented_string + '  -- ' + ql
    return commented_string[:-1]


def get_join_info(dataset, epa_table_name, wheresql, schema='public'):
    joins = []
    conn = connect_db()
    cur = conn.cursor() 
    sql = f"""select case when deagg_table_name is not null then deagg_table_name else organization_table_name end as organization_table_name, 
                    organization_join_table,
                    organization_join_column, organization_join_fk,
                    organization_join_column2, organization_join_fk2,
                    organization_join_column3, organization_join_fk3,
                    min(column_sort_order)
            from {schema}.v_{dataset.ust_or_release}_element_mapping_joins
            where {dataset.ust_or_release}_control_id = %s 
            and epa_table_name = %s """
    sql = sql + wheresql
    sql = sql + """\ngroup by case when deagg_table_name is not null then deagg_table_name else organization_table_name end, 
                    organization_join_table,
                    organization_join_column, organization_join_fk,
                    organization_join_column2, organization_join_fk2,
                    organization_join_column3, organization_join_fk3"""
    sql = sql + "\norder by 9"
    process_sql(conn, cur, sql, params=(dataset.control_id, epa_table_name))
    for row in cur.fetchall():
        join_info = {}
        join_info['organization_table_name'] = row[0]
        join_info['organization_join_table'] = row[1]
        join_info['organization_join_column'] = row[2]
        join_info['organization_join_fk'] = row[3]
        join_info['organization_join_column2'] = row[4]
        join_info['organization_join_fk2'] = row[5]
        join_info['organization_join_column3'] = row[6]
        join_info['organization_join_fk3'] = row[7]
        joins.append(join_info)
    cur.close()
    conn.close()    
    return joins


def get_lookup_info(dataset, epa_table_name, schema='public'):
    joins = []
    
    conn = connect_db()
    cur = conn.cursor() 
    sql = f"""select distinct database_lookup_column, organization_column_name 
            from {schema}.v_{dataset.ust_or_release}_element_mapping_joins
            where {dataset.ust_or_release}_control_id = %s and epa_table_name = %s
            and database_lookup_table is not null order by 1"""
    process_sql(conn, cur, sql, params=(dataset.control_id, epa_table_name))
    i = 1
    for row in cur.fetchall():
        join_info = {}
        join_info['organization_table_name'] = 'v_' + row[0] + '_xwalk' 
        join_info['organization_join_table'] = None
        join_info['organization_join_column'] = row[1]
        join_info['organization_join_fk'] = None
        join_info['organization_join_column2'] = None
        join_info['organization_join_fk2'] = None
        join_info['organization_join_column3'] = None
        join_info['organization_join_fk3'] = None   
        join_info['table_type'] = 'lookup'
        joins.append(join_info)
        i += 1
    cur.close()
    conn.close()     
    return joins


def get_join_tables(dataset, epa_table_name, schema='public'):
    aliases = list(string.ascii_lowercase)
    i = 0
    tables = []
    joins = []

    key_wheresql = " and primary_key is not null and organization_table_name not like 'erg_%%' "
    join_infos = get_join_info(dataset, epa_table_name, key_wheresql, schema)
    for join_info in join_infos:
        if join_info and join_info['organization_table_name'] not in tables:
            tables.append(join_info['organization_table_name'])
            join_info['table_type'] = 'key'
            join_info['alias'] = aliases[i]
            i += 1
            joins.append(join_info)

    org_wheresql = " and organization_table_name not like 'erg_%%' and database_lookup_column is null "
    join_infos = get_join_info(dataset, epa_table_name, org_wheresql, schema)
    for join_info in join_infos:
        if join_info and join_info['organization_table_name'] not in tables:
            tables.append(join_info['organization_table_name'])
            join_info['table_type'] = 'org'
            join_info['alias'] = aliases[i]
            i += 1
            joins.append(join_info)

    id_wheresql = " and organization_table_name like 'erg_%%' "
    join_infos = get_join_info(dataset, epa_table_name, id_wheresql, schema)
    for join_info in join_infos:
        if join_info and join_info['organization_table_name'] not in tables:
            tables.append(join_info['organization_table_name'])
            join_info['table_type'] = 'id'
            join_info['alias'] = aliases[i]
            i += 1
            joins.append(join_info)

    lookups = get_lookup_info(dataset, epa_table_name)
    if not tables and lookups:
        logger.warning('No base join table found for %s; skipping lookup joins.', epa_table_name)
        return joins
    for join_info in lookups:
        join_info['organization_join_table'] = tables[0]
        join_info['alias'] = aliases[i]
        i += 1
        joins.append(join_info)

    # TODO: add deagg joins  
    return joins 


def process_sql(conn, cur, sql, params=None, print_sql=False, exit_on_fail=True):
    try: 
        cur.execute(sql, params)
    except Exception as e:
        error_logger.error('Error processing SQL: %s', e, stack_info=True, exc_info=True)
        q = get_pretty_query(cur)
        error_logger.error('\n\nBad SQL:\n')
        error_logger.error(q)
        conn.rollback()
        if exit_on_fail:
            cur.close()
            conn.close()     
            error_logger.error('\n\nEXITING DUE TO SQL ERROR....\n\n')  
            raise RuntimeError('SQL execution failed; see logged SQL and database error details.') from e
    if print_sql:
        print(get_pretty_query(cur))


def pretty_print_df(df):
    print(df.to_markdown()) 


def string_to_list(string):
    if isinstance(string, list):
        return string
    if string == ', ':
        return[string]
    string = string.replace("'",'')
    str_list = string.split(', ')
    return str_list


def list_to_quoted_string(list_of_strings, delimiter=', '):
    """
    Takes a list of strings and returns a string with each list element quoted in single quotes.
    """
    if isinstance(list_of_strings, str):
        list_of_strings = string_to_list(list_of_strings)
    string = ''
    if list_of_strings and (type(list_of_strings) == list or type(list_of_strings) == tuple):
        string =  delimiter.join("'" + x + "'" for x in list_of_strings)
    elif list_of_strings:
        string = "'" + list_of_strings + "'"
    string = string.replace("''","'")
    if string == '':
        string = None
    return string    


def list_to_string(list_of_strings, delimiter=', '):
    """
    Takes a list of strings and returns a string without internal quotes.
    """    
    if isinstance(list_of_strings, str):
        return list_of_strings
    string = ''
    if list_of_strings and (type(list_of_strings) == list or type(list_of_strings) == tuple):
        string =  delimiter.join(x for x in list_of_strings)
    if string == '':
        string = None
    return string    


def list_nums_to_string(list_of_nums):
    if isinstance(list_of_nums, str):
        return list_of_nums
    string = ''
    if list_of_nums and (type(list_of_nums) == list or type(list_of_nums) == tuple):
        string =  ', '.join(str(x) for x in list_of_nums)
    elif list_of_nums:
        string = list_of_nums
    if string == '':
        string = None
    return string


def get_first_name_from_erg_email(email_address):
    i = email_address.find('.')
    return email_address[:i].title()


def get_last_name_from_erg_email(email_address):
    i = email_address.find('.')+1
    i2 = email_address.find('@')
    return email_address[i:i2].title()


def get_outlook_info():
    import win32com.client

    outlook = win32com.client.Dispatch('Outlook.Application')
    mapi = outlook.GetNameSpace('MAPI')
    outlook_info = {}
    outlook_info['email'] = mapi.Accounts[0].SmtpAddress
    outlook_info['first_name'] = get_first_name_from_erg_email(outlook_info['email'])
    outlook_info['last_name'] = get_last_name_from_erg_email(outlook_info['email'])
    outlook_info['display_name'] = mapi.Accounts[0].DisplayName
    return outlook_info
